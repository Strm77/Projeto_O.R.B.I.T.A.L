#!/usr/bin/env python3
"""Gera a planilha de trabalho (.xlsx) de um certame a partir do JSON da análise.

Abas:
  Checklist Habilitação       exigência | tipo | fonte | status | validade | obs
  Matriz Requisitos Técnicos  requisito | fonte | atende | evidência | responsável | obs
  Prazos                      evento | data/hora | fonte | dias restantes
  Go-No-Go                    critério | resultado | fonte   ("/" não é permitido em nome de aba)

Cabeçalho fixo, filtros, listas de validação e cores condicionais:
status/atende/resultado (verde, amarelo, vermelho) e prazos/validades com
≤ 3 dias em vermelho (≤ 7 em amarelo). "Dias restantes" é fórmula
(=data − HOJE()), então se atualiza sozinha ao abrir a planilha.

Uso:
    python gerar_planilhas.py analise.json -o planilhas.xlsx
    python gerar_planilhas.py --exemplo > analise.json   # modelo do JSON de entrada

Formato de entrada (todas as listas são opcionais):
{
  "certame": "Pregão Eletrônico nº 90042/2026 — AETI-VS",
  "habilitacao": [{"exigencia": "...", "tipo": "técnica", "fonte": "[edital.pdf, item 8.5.2, p. 3–4]",
                   "status": "Pendente", "validade": "2026-12-31", "obs": ""}],
  "requisitos":  [{"requisito": "...", "fonte": "...", "atende": "S", "evidencia": "",
                   "responsavel": "", "obs": ""}],
  "prazos":      [{"evento": "...", "data_hora": "2026-11-03 10:00", "fonte": "..."}],
  "go_no_go":    [{"criterio": "...", "resultado": "ATENÇÃO", "fonte": "..."}]
}
Datas: "aaaa-mm-dd[ hh:mm]" ou "dd/mm/aaaa[ hh:mm]". Prazo relativo sem data
(ex.: "convocação + 2 h") pode ir como texto em data_hora: fica sem contagem.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    sys.exit("Requer openpyxl: pip install openpyxl")

LINHAS_VALIDACAO = 500  # listas de validação valem também para linhas acrescentadas à mão

VERMELHO = (PatternFill("solid", bgColor="FFC7CE"), Font(color="9C0006"))
AMARELO = (PatternFill("solid", bgColor="FFEB9C"), Font(color="9C5700"))
VERDE = (PatternFill("solid", bgColor="C6EFCE"), Font(color="006100"))
CINZA = (PatternFill("solid", bgColor="EDEDED"), Font(color="595959"))

STATUS_HABILITACAO = {"Atende": VERDE, "Pendente": AMARELO, "Verificar": AMARELO,
                      "Não atende": VERMELHO, "N/A": CINZA}
ATENDE = {"S": VERDE, "Parcial": AMARELO, "N": VERMELHO}
RESULTADO = {"GO": VERDE, "ATENÇÃO": AMARELO, "NO-GO": VERMELHO, "NÃO ENCONTRADO": CINZA}
TIPOS = ["jurídica", "fiscal", "econômica", "técnica"]

# (chave no JSON, título da coluna, largura)
ABAS = {
    "Checklist Habilitação": ("habilitacao", [
        ("exigencia", "Exigência", 60), ("tipo", "Tipo", 12), ("fonte", "Fonte (doc/item/página)", 45),
        ("status", "Status", 13), ("validade", "Validade", 13), ("obs", "Obs", 40)]),
    "Matriz Requisitos Técnicos": ("requisitos", [
        ("requisito", "Requisito", 60), ("fonte", "Fonte", 45), ("atende", "Atende (S/N/Parcial)", 12),
        ("evidencia", "Evidência", 35), ("responsavel", "Responsável", 18), ("obs", "Obs", 35)]),
    "Prazos": ("prazos", [
        ("evento", "Evento", 55), ("data_hora", "Data/hora", 17), ("fonte", "Fonte", 50),
        ("dias", "Dias restantes", 10)]),
    "Go-No-Go": ("go_no_go", [
        ("criterio", "Critério", 60), ("resultado", "Resultado", 17), ("fonte", "Fonte", 55)]),
}

EXEMPLO = {
    "certame": "Pregão Eletrônico nº 90042/2026 — AETI-VS",
    "habilitacao": [{"exigencia": "Atestado: Service Desk SaaS ITSM, ≥150 agentes, ≥12 meses",
                     "tipo": "técnica", "fonte": "[edital.pdf, item 8.5.2, p. 3–4]",
                     "status": "Pendente", "validade": "", "obs": ""}],
    "requisitos": [{"requisito": "RF-01 Registrar chamados por portal, e-mail, app e chat (POC)",
                    "fonte": "[edital.pdf, Anexo II — Requisitos Técnicos da Solução, item 2.1, p. 9]",
                    "atende": "", "evidencia": "", "responsavel": "", "obs": ""}],
    "prazos": [{"evento": "Abertura da sessão pública", "data_hora": "2026-11-03 10:00",
                "fonte": "[errata.pdf, item 1, p. 1]"}],
    "go_no_go": [{"criterio": "B1 Exclusividade ME/EPP", "resultado": "GO",
                  "fonte": "[edital.pdf, item 2.2, p. 1]"}],
}


def _data(valor):
    """Converte texto de data em datetime; devolve o texto original se não for data."""
    if not valor or not isinstance(valor, str):
        return valor
    texto = re.sub(r"(\d{1,2})h(\d{2})", r"\1:\2", valor.strip())  # "10h00" → "10:00"
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M", "%Y-%m-%d", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            return datetime.strptime(texto, fmt)
        except ValueError:
            continue
    return valor


def _sem_emoji(texto):
    return re.sub(r"^[^\wÀ-ú]+", "", texto or "").strip()


def _cores_por_valor(ws, col, mapa, ultima):
    letra = get_column_letter(col)
    for valor, (fill, font) in mapa.items():
        ws.conditional_formatting.add(
            f"{letra}2:{letra}{ultima}",
            FormulaRule(formula=[f'${letra}2="{valor}"'], fill=fill, font=font, stopIfTrue=True))


def _validacao(ws, col, opcoes):
    letra = get_column_letter(col)
    dv = DataValidation(type="list", formula1='"' + ",".join(opcoes) + '"', allow_blank=True,
                        showErrorMessage=False)  # aceita outro valor, só sugere a lista
    dv.add(f"{letra}2:{letra}{LINHAS_VALIDACAO}")
    ws.add_data_validation(dv)


def _aba(wb, titulo, colunas, linhas):
    ws = wb.create_sheet(titulo)
    cab_fill = PatternFill("solid", fgColor="1F3864")
    for c, (_, nome, largura) in enumerate(colunas, start=1):
        cel = ws.cell(row=1, column=c, value=nome)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = cab_fill
        cel.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = largura
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    for r, linha in enumerate(linhas, start=2):
        for c, (chave, _, _) in enumerate(colunas, start=1):
            if chave == "dias":
                valor = f'=IF(ISNUMBER(B{r}),INT(B{r})-TODAY(),"")'
            elif chave in ("data_hora", "validade"):
                valor = _data(linha.get(chave))
            elif chave in ("status", "atende", "resultado"):
                valor = _sem_emoji(linha.get(chave))
            else:
                valor = linha.get(chave)
            cel = ws.cell(row=r, column=c, value=valor)
            cel.alignment = Alignment(vertical="top", wrap_text=True)
            if chave == "dias":
                cel.number_format = "0"  # sem isso o Calc/Excel mostra a diferença como data
                cel.alignment = Alignment(vertical="top", horizontal="center")
            if isinstance(valor, datetime):
                cel.number_format = "dd/mm/yyyy hh:mm" if (valor.hour or valor.minute) else "dd/mm/yyyy"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    ultima = max(len(linhas) + 1, LINHAS_VALIDACAO)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{len(linhas) + 1}"
    return ws, ultima


def gerar(dados: dict, destino: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = dados.get("certame") or "O.R.B.I.T.A.L"
    wb.properties.creator = "orbital/gerar_planilhas.py"

    for titulo, (chave, colunas) in ABAS.items():
        ws, ultima = _aba(wb, titulo, colunas, dados.get(chave) or [])
        if chave == "habilitacao":
            _validacao(ws, 2, TIPOS)
            _validacao(ws, 4, list(STATUS_HABILITACAO))
            _cores_por_valor(ws, 4, STATUS_HABILITACAO, ultima)
            for limite, (fill, font) in ((3, VERMELHO), (7, AMARELO)):
                ws.conditional_formatting.add(f"E2:E{ultima}", FormulaRule(
                    formula=[f"AND(ISNUMBER($E2),INT($E2)-TODAY()<={limite})"],
                    fill=fill, font=font, stopIfTrue=True))
        elif chave == "requisitos":
            _validacao(ws, 3, list(ATENDE))
            _cores_por_valor(ws, 3, ATENDE, ultima)
        elif chave == "prazos":
            for limite, (fill, font) in ((3, VERMELHO), (7, AMARELO)):
                ws.conditional_formatting.add(f"A2:D{ultima}", FormulaRule(
                    formula=[f"AND(ISNUMBER($D2),$D2<={limite})"], fill=fill, font=font, stopIfTrue=True))
        elif chave == "go_no_go":
            _validacao(ws, 2, list(RESULTADO))
            _cores_por_valor(ws, 2, RESULTADO, ultima)

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("analise", nargs="?", type=Path, help="JSON da análise")
    ap.add_argument("-o", "--saida", type=Path, default=Path("planilhas.xlsx"))
    ap.add_argument("--exemplo", action="store_true", help="imprime um JSON de entrada de exemplo")
    args = ap.parse_args(argv)

    if args.exemplo:
        print(json.dumps(EXEMPLO, ensure_ascii=False, indent=2))
        return 0
    if not args.analise or not args.analise.is_file():
        print("erro: informe o JSON da análise (ou use --exemplo)", file=sys.stderr)
        return 2
    dados = json.loads(args.analise.read_text(encoding="utf-8"))
    destino = gerar(dados, args.saida)
    contagem = {t: len(dados.get(c) or []) for t, (c, _) in ABAS.items()}
    print(f"{destino}: " + ", ".join(f"{t}={n}" for t, n in contagem.items()), file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
