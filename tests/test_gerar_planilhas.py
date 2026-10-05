"""Testes do gerar_planilhas.py."""

import json
from datetime import datetime

from openpyxl import load_workbook

import gerar_planilhas as gp

ABAS = ["Checklist Habilitação", "Matriz Requisitos Técnicos", "Prazos", "Go-No-Go"]


def gerar(tmp_path, dados=None):
    destino = tmp_path / "p.xlsx"
    gp.gerar(dados or gp.EXEMPLO, destino)
    return load_workbook(destino)


def test_abas_e_cabecalhos(tmp_path):
    wb = gerar(tmp_path)
    assert wb.sheetnames == ABAS
    assert [c.value for c in wb["Checklist Habilitação"][1]] == [
        "Exigência", "Tipo", "Fonte (doc/item/página)", "Status", "Validade", "Obs"]
    assert [c.value for c in wb["Matriz Requisitos Técnicos"][1]] == [
        "Requisito", "Fonte", "Atende (S/N/Parcial)", "Evidência", "Responsável", "Obs"]
    assert [c.value for c in wb["Prazos"][1]] == ["Evento", "Data/hora", "Fonte", "Dias restantes"]
    assert [c.value for c in wb["Go-No-Go"][1]] == ["Critério", "Resultado", "Fonte"]


def test_cabecalho_fixo_e_filtros(tmp_path):
    wb = gerar(tmp_path)
    for ws in wb:
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref and ws.auto_filter.ref.startswith("A1:")


def test_prazos_data_formula_e_cor(tmp_path):
    dados = {"prazos": [
        {"evento": "Sessão", "data_hora": "03/11/2026 10h00", "fonte": "[x.pdf, item 1, p. 1]"},
        {"evento": "Proposta ajustada", "data_hora": "convocação + 2 h", "fonte": "[x.pdf, item 5.3, p. 2]"},
    ]}
    ws = gerar(tmp_path, dados)["Prazos"]
    assert ws["B2"].value == datetime(2026, 11, 3, 10, 0)
    assert ws["B3"].value == "convocação + 2 h"  # relativo fica como texto, sem contagem
    assert ws["D2"].value == '=IF(ISNUMBER(B2),INT(B2)-TODAY(),"")'
    assert ws["D2"].number_format == "0"
    formulas = [r.formula[0] for faixa in ws.conditional_formatting for r in faixa.rules]
    assert "AND(ISNUMBER($D2),$D2<=3)" in formulas
    vermelho = next(r for faixa in ws.conditional_formatting for r in faixa.rules
                    if r.formula[0] == "AND(ISNUMBER($D2),$D2<=3)")
    assert vermelho.dxf.fill.bgColor.rgb.endswith("FFC7CE")


def test_cores_de_status(tmp_path):
    wb = gerar(tmp_path)
    regras = {r.formula[0] for faixa in wb["Checklist Habilitação"].conditional_formatting for r in faixa.rules}
    assert {'$D2="Atende"', '$D2="Não atende"', '$D2="Pendente"'} <= regras
    regras = {r.formula[0] for faixa in wb["Matriz Requisitos Técnicos"].conditional_formatting for r in faixa.rules}
    assert {'$C2="S"', '$C2="N"', '$C2="Parcial"'} <= regras
    regras = {r.formula[0] for faixa in wb["Go-No-Go"].conditional_formatting for r in faixa.rules}
    assert {'$B2="GO"', '$B2="NO-GO"', '$B2="ATENÇÃO"'} <= regras


def test_resultado_com_emoji_e_normalizado(tmp_path):
    ws = gerar(tmp_path, {"go_no_go": [{"criterio": "x", "resultado": "⛔ NO-GO", "fonte": ""}]})["Go-No-Go"]
    assert ws["B2"].value == "NO-GO"


def test_cli(tmp_path):
    entrada = tmp_path / "a.json"
    entrada.write_text(json.dumps(gp.EXEMPLO, ensure_ascii=False), encoding="utf-8")
    saida = tmp_path / "out.xlsx"
    assert gp.main([str(entrada), "-o", str(saida)]) == 0
    assert saida.is_file()
    assert gp.main([str(tmp_path / "nao_existe.json")]) == 2
